import io
import json
import os
from pathlib import Path
import tempfile
import time
from types import SimpleNamespace
import unittest
from unittest.mock import patch

from PIL import Image
from docx import Document
from starlette.testclient import TestClient
from backend.agent_core import AgentCore
from backend.api import create_app
from backend.uploaded_files import prepare_file
from test_backend import Message, MockClient


def fixtures():
    document = Document(); document.add_paragraph('Orion project: Python; 80ms to 20ms.')
    table = document.add_table(rows=1, cols=2); table.cell(0, 0).text = 'Table fact'; table.cell(0, 1).text = 'kept'
    docx = io.BytesIO(); document.save(docx)
    image = io.BytesIO(); Image.new('RGB', (60, 40), 'blue').save(image, format='PNG')
    return docx.getvalue(), image.getvalue()


class UploadClient(MockClient):
    observed = []
    reject_once = False
    async def create(self, **kwargs):
        self.calls += 1
        self.observed.append(kwargs['messages'])
        if type(self).reject_once:
            type(self).reject_once = False
            raise RuntimeError('private failure')
        if self.calls == 1:
            material = json.loads(kwargs['messages'][1]['content'].split('\n',1)[1])['markdown']
            message = Message(tool_calls=[SimpleNamespace(id='save', function=SimpleNamespace(
                name='resume__write_file', arguments=json.dumps({'path':'项目经历库.md','content':material+'\n- Orion project: Python; 80ms to 20ms.\n'})))])
        else:
            message = Message(content='已读取文字和附件，并保存确认的项目。')
        return SimpleNamespace(choices=[SimpleNamespace(message=message)])


def core_factory(store):
    return AgentCore(store, client_factory=UploadClient)


def wait_task(client, base, tid, headers):
    for _ in range(300):
        result = client.get(base+'/tasks/'+tid, headers=headers).json()
        if result['status'] in {'completed','failed'}:
            return result
        time.sleep(.02)
    raise AssertionError('task did not finish')


class UploadTests(unittest.TestCase):
    def test_raw_upload_model_images_history_ownership_and_restart(self):
        docx, png = fixtures(); UploadClient.observed = []
        with tempfile.TemporaryDirectory() as directory, patch.dict(os.environ, AGENT_API_KEY='test', AGENT_MODEL='test'):
            with TestClient(create_app(directory, core_factory, allow_legacy_users=True)) as client:
                a,b = client.post('/users').json(),client.post('/users').json()
                ha,hb = ({'Authorization':'Bearer '+u['token']} for u in (a,b))
                sid = client.post('/sessions',headers=ha).json()['session_id'];base='/sessions/'+sid
                result = client.post(base+'/messages',headers=ha,data={'message':'同时读取我的项目文档和图片，不要忽略本句要求'},files=[
                    ('files',('项目.docx',docx,'application/octet-stream')),('files',('截图.png',png,'image/png'))])
                self.assertEqual(result.status_code,202,result.text)
                task = wait_task(client,base,result.json()['task_id'],ha)
                self.assertEqual(task['status'],'completed',task)
                supplied = next(m for m in UploadClient.observed[0] if m['role']=='user' and isinstance(m['content'],list))['content']
                self.assertIn('不要忽略',supplied[0]['text'])
                self.assertIn('Orion project',json.dumps(supplied))
                self.assertIn('Table fact',json.dumps(supplied))
                self.assertEqual(len([p for p in supplied if p['type']=='image_url']),1)
                history = client.get(base+'/messages',headers=ha).json()['messages']
                self.assertEqual(len(history[0]['attachments']),2)
                attachment = history[0]['attachments'][0]
                path = base+'/uploads/'+attachment['id']
                self.assertEqual(client.get(path,headers=ha).content,docx)
                self.assertEqual(client.get(path,headers=hb).status_code,404)
                self.assertEqual(client.get(path).status_code,401)
                sid2=client.post('/sessions',headers=ha).json()['session_id']
                self.assertEqual(client.get('/sessions/'+sid2+'/uploads/'+attachment['id'],headers=ha).status_code,404)
                self.assertIn('Orion',client.get('/materials',headers=ha).json()['markdown'])
                stored=client.app.state.store.history(a['user_id'],sid)
                self.assertNotIn('data:image',json.dumps(stored))
            with TestClient(create_app(directory,core_factory, allow_legacy_users=True)) as client:
                self.assertEqual(client.get(path,headers=ha).content,docx)
                tid=client.post(base+'/messages',headers=ha,json={'message':'继续看看上次图片'}).json()['task_id']
                self.assertEqual(wait_task(client,base,tid,ha)['status'],'completed')
                self.assertTrue(any(isinstance(m.get('content'),list) for m in UploadClient.observed[-1]))

    def test_validation_cleanup_and_retry_preserves_raw_files(self):
        _,png=fixtures()
        with tempfile.TemporaryDirectory() as directory,patch.dict(os.environ,AGENT_API_KEY='test',AGENT_MODEL='test'):
            with TestClient(create_app(directory,core_factory, allow_legacy_users=True)) as client:
                u=client.post('/users').json();h={'Authorization':'Bearer '+u['token']}
                sid=client.post('/sessions',headers=h).json()['session_id'];base='/sessions/'+sid
                bad=client.post(base+'/messages',headers=h,files=[('files',('ok.png',png)),('files',('bad.exe',b'bad'))])
                self.assertEqual(bad.status_code,400)
                with client.app.state.store.db() as db:self.assertEqual(db.execute('SELECT count(*) FROM uploads').fetchone()[0],0)
                too_many=client.post(base+'/messages',headers=h,files=[('files',(str(i)+'.png',png)) for i in range(4)])
                self.assertEqual(too_many.status_code,400)
                corrupt=client.post(base+'/messages',headers=h,files=[('files',('bad.png',b'not an image'))]).json()['task_id']
                self.assertEqual(wait_task(client,base,corrupt,h)['status'],'failed')
                self.assertEqual(client.get(base+'/messages',headers=h).json()['messages'],[])
                UploadClient.reject_once=True
                tid=client.post(base+'/messages',headers=h,data={'message':'读取图片'},files=[('files',('image.png',png))]).json()['task_id']
                self.assertEqual(wait_task(client,base,tid,h)['status'],'failed')
                retry=client.post(base+'/tasks/'+tid+'/retry',headers=h).json()['task_id']
                self.assertEqual(wait_task(client,base,retry,h)['status'],'completed')
                history=client.get(base+'/messages',headers=h).json()['messages']
                self.assertEqual(len(history[0]['attachments']),1)

    def test_scanned_pdf_is_sent_as_pages_and_password_is_rejected(self):
        from pypdf import PdfWriter
        with tempfile.TemporaryDirectory() as directory:
            root=Path(directory);pdf=root/'scan.pdf'
            Image.new('RGB',(250,350),'white').save(pdf,format='PDF')
            data=prepare_file(pdf,root/'ready')
            self.assertFalse(data['text'].strip());self.assertEqual(len(data['images']),1)
            protected=root/'protected.pdf';writer=PdfWriter();writer.add_blank_page(width=200,height=300);writer.encrypt('secret')
            with protected.open('wb') as f:writer.write(f)
            with self.assertRaisesRegex(ValueError,'密码'):
                prepare_file(protected,root/'locked')


if __name__=='__main__':unittest.main()
